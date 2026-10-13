"""画像面试回归：独立自我介绍、合同验证、提示、追问轮次与失败恢复。"""
import base64
import io
import json
import zipfile

import pytest
from langgraph.checkpoint.memory import InMemorySaver
from langgraph.types import Command

from app.interview.workflow import build_interview_graph
from app.interview.setup import catalog, extract_text
from app.model import ModelFailure


def initial(total=2):
    return {"setup": {"workflowVersion": "v2", "directionId": "java_backend", "level": "junior",
                      "depth": "short", "maxQuestions": total, "allowFollowup": True,
                      "candidateProfile": {"summary": "应届生", "skills": ["Java"], "projects": [], "concerns": []},
                      "jobProfile": {"title": "Java", "skills": ["Java"], "requirements": [], "dimensions": ["technical_depth"]}},
            "questions": [{"id": "1", "title": "线程池", "content": "作用", "referenceAnswer": "复用线程", "sourceVersion": "v1"}],
            "main_index": 0, "turns": [], "status": "preparing"}


class Model:
    """根据节点 Schema 返回结构化样例，错误证据只触发一次以验收恢复。"""
    def __init__(self, invalid_evidence=False):
        self.calls = 0
        self.invalid = invalid_evidence
        self.ask_count = 0
    def __call__(self, messages):
        self.calls += 1
        prompt = messages[0]["content"]
        data = json.loads(messages[1]["content"])
        if "画像" in prompt and "自我介绍" in data:
            return json.dumps(data["已有画像"], ensure_ascii=False)
        if "面试导演" in prompt:
            self.ask_count += 1
            return json.dumps({"title": "线程池", "content": "为什么使用线程池？", "dimension": "technical_depth",
                "sourceId": "1" if self.ask_count == 1 else "", "projectIndex": None,
                "mustCover": ["线程复用", "并发控制"], "minimumBar": "解释核心作用", "reason": "岗位技能匹配"}, ensure_ascii=False)
        if "引导性问题" in prompt:
            return json.dumps({"clue": "想一想线程创建成本。"}, ensure_ascii=False)
        if "已冻结合同" in prompt:
            evidence = ["凭空引用"] if self.invalid else ["复用线程"]
            self.invalid = False
            return json.dumps({"accuracy": 80, "completeness": 60, "expression": 70,
                "dimensionScore": 70, "covered": [0], "missing": [1], "intent": "substantive",
                "strengths": ["提到复用"], "improvements": ["补充并发控制"], "evidence": evidence,
                "followup": "如何控制并发？"}, ensure_ascii=False)
        return json.dumps({"summary": "理解复用，需要补充并发控制。", "trainingPlan": [{"goal": "理解并发控制",
            "action": "说明线程池容量与队列的关系", "criterion": "能解释拒绝策略", "turnId": "turn-1"}]}, ensure_ascii=False)


def test_intro_hint_and_followup_share_formal_budget():
    model = Model()
    saver = InMemorySaver()
    graph = build_interview_graph(saver, model, "v2")
    config = {"configurable": {"thread_id": "v2"}}
    graph.invoke(initial(), config)
    assert model.calls == 0
    assert graph.get_state(config).values["current"]["kind"] == "self_intro"
    graph.invoke(Command(resume={"action": "answer", "answer": "我是应届生，学习 Java。"}), config)
    state = graph.get_state(config).values
    assert state["main_index"] == 0 and state["turns"][0]["evaluation"] is None
    assert "mustCover" not in state["current"] and state["contract"]["mustCover"]
    graph.invoke(Command(resume={"action": "hint"}), config)
    state = graph.get_state(config).values
    assert state["main_index"] == 0 and len(state["current"]["hints"]) == 1
    graph.invoke(Command(resume={"action": "answer", "answer": "复用线程"}), config)
    assert graph.get_state(config).values["current"]["kind"] == "followup"
    # 使用相同检查点构建新实例，模拟重启后继续，而不是重新创建面试。
    graph = build_interview_graph(saver, model, "v2")
    graph.invoke(Command(resume={"action": "answer", "answer": "复用线程，同时控制并发。"}), config)
    state = graph.get_state(config).values
    assert state["status"] == "completed" and state["main_index"] == 2
    assert state["report"]["trainingStatus"] == "succeeded"
    assert state["report"]["answered"] == 2
    assert next(d for d in state["report"]["dimensionScores"] if d["id"] == "system_design")["score"] is None


def test_invalid_citation_does_not_advance_and_retry_keeps_original_answer():
    model = Model(invalid_evidence=True)
    graph = build_interview_graph(InMemorySaver(), model, "v2")
    config = {"configurable": {"thread_id": "citation"}}
    graph.invoke(initial(1), config)
    graph.invoke(Command(resume={"action": "skip"}), config)
    with pytest.raises(ModelFailure, match="EVALUATION_INVALID"):
        graph.invoke(Command(resume={"action": "answer", "answer": "复用线程\n原回答"}), config)
    assert graph.get_state(config).values["main_index"] == 0
    # 从评价节点重跑才能修复错误输出，验证失败保留原始答案。
    graph.update_state(config, {"action": "answer"}, as_node="wait_answer")
    graph.invoke(None, config)
    assert graph.get_state(config).values["turns"][-1]["answer"] == "复用线程\n原回答"


def test_early_end_has_no_fabricated_score_or_model_calls():
    graph = build_interview_graph(InMemorySaver(), lambda messages: pytest.fail("No completed answer"), "v2")
    config = {"configurable": {"thread_id": "end"}}
    graph.invoke(initial(), config)
    graph.invoke(Command(resume={"action": "end"}), config)
    state = graph.get_state(config).values
    assert state["report"]["score"] is None and state["report"]["trainingPlan"] == []


def test_resume_text_and_docx_tables_are_extracted():
    from docx import Document
    document = Document()
    document.add_paragraph("应届生 Java")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "面试蛙项目"
    stream = io.BytesIO(); document.save(stream)
    result = extract_text({"kind": "resume", "filename": "resume.docx", "fileBase64": base64.b64encode(stream.getvalue()).decode()})
    assert "面试蛙项目" in result and "Java" in result


@pytest.mark.parametrize("payload", [
    {"kind": "resume", "filename": "fake.pdf", "fileBase64": base64.b64encode(b"not pdf").decode()},
    {"kind": "resume", "text": ""},
    {"kind": "jd", "text": "a" * 12001},
])
def test_invalid_documents_fail_explicitly(payload):
    with pytest.raises(ValueError):
        extract_text(payload)


def test_direction_catalog_is_upstream_data_and_depths_are_5_8_12():
    result = catalog()
    assert len(result["directions"]) > 3
    assert [d["turns"] for d in result["depths"]] == [5, 8, 12]


def test_hint_rejects_overlong_answer_instead_of_showing_it():
    class OverlongHint(Model):
        def __call__(self, messages):
            if "引导性问题" in messages[0]["content"]:
                return json.dumps({"clue": "完整答案" * 100}, ensure_ascii=False)
            return super().__call__(messages)
    graph = build_interview_graph(InMemorySaver(), OverlongHint(), "v2")
    config = {"configurable": {"thread_id": "hint-boundary"}}
    graph.invoke(initial(), config)
    graph.invoke(Command(resume={"action": "skip"}), config)
    with pytest.raises(ModelFailure, match="SETUP_OUTPUT_INVALID"):
        graph.invoke(Command(resume={"action": "hint"}), config)
    assert graph.get_state(config).values["current"]["hints"] == []
    assert graph.get_state(config).values["main_index"] == 0


def test_generated_question_cannot_override_confirmed_dimensions():
    class WrongDimension(Model):
        def __call__(self, messages):
            reply = super().__call__(messages)
            if "面试导演" in messages[0]["content"]:
                data = json.loads(reply); data["dimension"] = "system_design"
                return json.dumps(data, ensure_ascii=False)
            return reply
    graph = build_interview_graph(InMemorySaver(), WrongDimension(), "v2")
    config = {"configurable": {"thread_id": "dimension-boundary"}}
    graph.invoke(initial(), config)
    with pytest.raises(ModelFailure, match="QUESTION_PLAN_INVALID"):
        graph.invoke(Command(resume={"action": "skip"}), config)


def test_retrieval_sources_recorded_in_private_checkpoint():
    seen = []
    def retrieve(questions, query, used):
        seen.append((query, used))
        return questions, {'mode': 'hybrid', 'hits': [{'questionId': '1', 'sourceVersion': 'v1'}]}
    graph = build_interview_graph(InMemorySaver(), Model(), 'v2', retrieve=retrieve)
    config = {'configurable': {'thread_id': 'retrieval'}}
    value = initial(1)
    value['setup']['jobPosition'] = 'Java后端'
    graph.invoke(value, config)
    graph.invoke(Command(resume={'action': 'skip'}), config)
    state = graph.get_state(config).values
    assert seen and 'Java' in seen[0][0]
    assert state['retrieval_archive']['turn-1']['mode'] == 'hybrid'
    assert state['contract_archive']['turn-1-contract']['referenceKnowledge'] == '复用线程'
    assert 'referenceKnowledge' not in state['current']
